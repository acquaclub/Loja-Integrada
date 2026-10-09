import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
import curva as c

ARQ = c.B + 'Pente-fino-SKUs-Unikitchen.xlsx'
F = 'Arial'
cand = c.cand.copy()

grupos = []
for gk, x in cand.groupby('grupo', sort=False):
    x = x.sort_values('pontos', ascending=False)
    r = x.iloc[0]
    precos = [p for p in x.preco if pd.notna(p) and p]
    grupos.append(dict(
        remessa='Sim' if r.remessa else ('Cortado' if all(x.corte != '') else 'Reserva'),
        diferenciado=next((d for d in x.diferenciado if d), ''),
        corte=next((d for d in x.corte if d), ''), categoria=r.cat_curta,
        marca=r.marca, curva=min(x.curva), pontos=r.pontos,
        codigos=', '.join(x.codigo), n=len(x),
        produto=r.descricao[:140], colecao=r.colecao,
        situacao=' / '.join(dict.fromkeys(x.situacao))[:120],
        preco=(min(precos) if precos else None),
        motivo=r.motivo, fonte=r.fonte, marca_site=r.marca_site))
gdf = pd.DataFrame(grupos)
gdf['ord_r'] = gdf.remessa.map({'Sim': 0, 'Reserva': 1, 'Cortado': 2})
gdf = gdf.sort_values(['ord_r', 'marca_site', 'marca', 'categoria', 'curva', 'pontos'], ascending=[True, False, True, True, True, False])

wb = load_workbook(ARQ)
for nome in ('Cadastro novos', 'Capacidade'):
    if nome in wb.sheetnames:
        del wb[nome]

# ---------- Capacidade ----------
cp = wb.create_sheet('Capacidade', 1)
sel = gdf[gdf.remessa == 'Sim']
linhas = [
    ('Limite do plano (produtos ativos)', c.LIMITE, 'Painel da Loja Integrada'),
    ('Ativos hoje', c.ATIVOS_HOJE, 'Painel: 196 de 500, menos 8 refrigeradores Smeg excluídos. Inativos e variações não contam'),
    ('Margem para marcas novas', c.MARGEM, 'Pedido do cliente'),
    ('Disponível (sem a margem)', '=B2-B3-B4', ''),
    ('Reativações sugeridas (aba Revisão do catálogo)', c.REATIVAR, 'Também passam a contar como ativos'),
    ('Vagas guardadas para as listas que faltam', c.RESERVA_LISTAS, 'Tramontina, Franke, InSinkErator, Deca, Jacuzzi, Speed Queen'),
    ('Remessa sugerida (anúncios)', len(sel), 'Aba "Cadastro novos", coluna Remessa = Sim'),
    ('Sobra', '=B5-B6-B7-B8', 'Deve dar zero'),
]
cp['A1'] = 'Capacidade de cadastro'
cp['A1'].font = Font(name=F, size=14, bold=True)
for i, (a, b, n) in enumerate(linhas, 2):
    cp.cell(row=i, column=1, value=a).font = Font(name=F, size=10, bold=(i in (5, 8, 9)))
    cp.cell(row=i, column=2, value=b).font = Font(name=F, size=10, bold=(i in (5, 8)), color='0000FF' if isinstance(b, int) and i in (2, 4) else '000000')
    cp.cell(row=i, column=3, value=n).font = Font(name=F, size=9, italic=True, color='666666')
notas = [
    'Atenção:',
    '• Ainda faltam as listas de Tramontina, Franke, InSinkErator, Deca, Jacuzzi e Speed Queen: já há 35 vagas guardadas para elas.',
    '• Indisponível com previsão conta como disponível: o site é catálogo, o cliente compra e aguarda.',
    '• Leque de produtos: dentro de cada marca a seleção passa por todas as categorias antes de repetir uma. Tetos: Tecno 30, Elettromec 30, Bertazzoni 35.',
    '• Os 26 anúncios da lista "Atualizar" (aba Plano de ação) não somam: trocam o código de anúncios que já existem.',
]
for k, t in enumerate(notas, 11):
    cp.cell(row=k, column=1, value=t).font = Font(name=F, size=10, bold=(k == 11))
cp.column_dimensions['A'].width = 62; cp.column_dimensions['B'].width = 12; cp.column_dimensions['C'].width = 80

# ---------- Cadastro novos ----------
ws = wb.create_sheet('Cadastro novos', 2)
tit = ['Remessa', 'Marca', 'Categoria', 'Curva', 'Código(s) na lista da marca', 'Variações', 'Produto (descrição da marca)', 'Diferenciado',
       'Situação na lista', 'Preço sugerido (R$)', 'Por que essa curva', 'Motivo do corte', 'Fonte', 'Cadastrado?']
larg = [10, 13, 16, 7, 34, 10, 66, 34, 30, 13, 46, 38, 24, 12]
for j, (t, w) in enumerate(zip(tit, larg), 1):
    cl = ws.cell(row=1, column=j, value=t)
    cl.font = Font(name=F, size=10, bold=True, color='FFFFFF'); cl.fill = PatternFill('solid', fgColor='1A1A1A')
    cl.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[cl.column_letter].width = w
ws.row_dimensions[1].height = 30
borda = Border(bottom=Side(style='thin', color='D9D9D9'))
for i, (_, r) in enumerate(gdf.iterrows(), 2):
    vals = [r.remessa, r.marca, r.categoria, r.curva, r.codigos, r.n, r.produto, r.diferenciado, r.situacao, r.preco, r.motivo, r.corte, r.fonte, None]
    for j, v in enumerate(vals, 1):
        cl = ws.cell(row=i, column=j, value=v)
        cl.font = Font(name=F, size=10); cl.border = borda
        cl.alignment = Alignment(vertical='top', wrap_text=j in (5, 7, 8, 11, 12))
    ws.cell(row=i, column=10).number_format = '#,##0'
ult = len(gdf) + 1
ws.freeze_panes = 'E2'
ws.auto_filter.ref = f'A1:N{ult}'
ws.conditional_formatting.add(f'D2:D{ult}', FormulaRule(formula=['$D2="A"'], fill=PatternFill('solid', fgColor='D4EDDA')))
ws.conditional_formatting.add(f'D2:D{ult}', FormulaRule(formula=['$D2="B"'], fill=PatternFill('solid', fgColor='FFF2CC')))
ws.conditional_formatting.add(f'D2:D{ult}', FormulaRule(formula=['$D2="C"'], fill=PatternFill('solid', fgColor='EEEEEE')))
ws.conditional_formatting.add(f'A2:N{ult}', FormulaRule(formula=['$A2<>"Sim"'], font=Font(name=F, size=10, color='888888')))
ws.conditional_formatting.add(f'H2:H{ult}', FormulaRule(formula=['$H2<>""'], fill=PatternFill('solid', fgColor='E8DCC4')))

# resumo por marca na aba Capacidade
cp.cell(row=17, column=1, value='Remessa por marca').font = Font(name=F, size=12, bold=True)
hdr = ['Marca', 'Anúncios na remessa', 'Diferenciados', 'Curva A', 'Curva B', 'Reserva (fica para depois)']
for j, t in enumerate(hdr, 1):
    cl = cp.cell(row=18, column=j, value=t); cl.font = Font(name=F, size=10, bold=True, color='FFFFFF'); cl.fill = PatternFill('solid', fgColor='1A1A1A')
for col, w in zip('DEF', [10, 10, 24]):
    cp.column_dimensions[col].width = w
marcas = list(dict.fromkeys(gdf.marca))
for i, m in enumerate(marcas, 19):
    R = f"'Cadastro novos'!"
    vals = [m,
            f'=COUNTIFS({R}$B$2:$B${ult},A{i},{R}$A$2:$A${ult},"Sim")',
            f'=COUNTIFS({R}$B$2:$B${ult},A{i},{R}$A$2:$A${ult},"Sim",{R}$H$2:$H${ult},"<>")',
            f'=COUNTIFS({R}$B$2:$B${ult},A{i},{R}$A$2:$A${ult},"Sim",{R}$D$2:$D${ult},"A")',
            f'=COUNTIFS({R}$B$2:$B${ult},A{i},{R}$A$2:$A${ult},"Sim",{R}$D$2:$D${ult},"B")',
            f'=COUNTIFS({R}$B$2:$B${ult},A{i},{R}$A$2:$A${ult},"Reserva")']
    for j, v in enumerate(vals, 1):
        cp.cell(row=i, column=j, value=v).font = Font(name=F, size=10)
t = 19 + len(marcas)
cp.cell(row=t, column=1, value='Total').font = Font(name=F, size=10, bold=True)
for j, col in enumerate('BCDEF', 2):
    cp.cell(row=t, column=j, value=f'=SUM({col}19:{col}{t-1})').font = Font(name=F, size=10, bold=True)

wb.save(ARQ)
print('anuncios remessa', len(sel), 'reserva', (gdf.remessa == 'Reserva').sum())
print(sel.groupby('marca').agg(anuncios=('n', 'size'), A=('curva', lambda s: (s == 'A').sum()), B=('curva', lambda s: (s == 'B').sum())))
