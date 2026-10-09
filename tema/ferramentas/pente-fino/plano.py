from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
F='Arial'
atual = [
 # (SKU site, produto, marca, substituto na lista, situação, observação)
 ('EFNEB0352A','Forno Elétrico Evol IX13 73L','Evol','EFNEB0352C ou EFNEB0352D','Disponível','Código mudou (já estava na lista de atualização)'),
 ('EFNEB0352B','Forno Elétrico Evol IX17 73L','Evol','EFNEB0352C ou EFNEB0352D','Disponível','Código mudou'),
 ('EMCEB0352A','Micro-ondas Combinado Evol IX13 35L','Evol','EMCEB0352B ou EMCEB0352C','Disponível','Código mudou'),
 ('EGAEB0062A','Gaveta Aquecida Evol IX17','Evol','EGAEB0062C ou EGAEB0062D','Disponível','Código mudou'),
 ('EGAEB0062B','Gaveta Aquecida Evol IX13','Evol','EGAEB0062C ou EGAEB0062D','Disponível','Código mudou'),
 ('ECTGD4QM3A','Cooktop a Gás Evol Inox 4 Bocas','Evol','ECTGD4QM3C ou ECTGD4QM3D','Disponível / em breve','Código mudou'),
 ('CBB300C','Churrasqueira Evol Modena 28"','Evol','CBY300','Disponível','Mesmo modelo com código novo'),
 ('JC145CL','Freezer Evol 142L','Evol','JC-145CL220R','Disponível','Código mudou'),
 ('TR08HC','Umidor Tecno Professional 38cm','Tecno','TR06HC (nova versão)','Disponível','Nova versão'),
 ('TG14EXAP','Gaveta Aquecida Tecno Professional','Tecno','TG14EXDP (Professional) / TG14EXDA (Original)','Indisponível – 2ª quinz. out','Modelo novo (digital)'),
 ('TO73EXDB','Forno Elétrico Tecno 73L','Tecno','TO81EXDAP (Professional) / TO81EXDA (Original)','Indisponível – 2ª quinz. out','Modelo novo de 81L'),
 ('TR14GXDA','Gaveta Refrigerada Tecno 136L','Tecno','TR10GFXDA (95L, com freezer)','Disponível','Modelo novo, medida diferente'),
 ('FR-LL-45-XX-NCLA','Frente Lava-louças 10 Serviços','Elettromec','FR-LL-45-XX-NPLA ou FR-LL-45-XX-NGZA','Em linha – chega out/dez','Código mudou'),
 ('FR-LL-60-XX-NCLB','Frente Lava-louças 14 Serviços','Elettromec','FR-LL-60-XX-NPLA ou FR-LL-60-XX-NGZA','Em linha – chega out/dez','Código mudou'),
 ('FR-LL-60-VT-NLGB','Painel Lava-louças Vetro 14 Serviços','Elettromec','FR-LL-60-XX-NPLA ou FR-LL-60-XX-NGZA (inox)','Em linha','Vetro saiu; só inox'),
 ('CKI-4Q-60-CI-2XBB','Cooktop Indução 4 Bocas 60cm','Elettromec','CKI-4Q-60-CI-2KSA','Em linha – disponível','Código mudou'),
 ('DI-2Q-30-CI-2XBB','Dominó Indução 2 Bocas 30cm','Elettromec','DI-2Q-30-CI-2KSA','Em linha – disponível','Código mudou'),
 ('FG-AN-90-LC-2LCA','Forno a Gás Luce 90cm','Elettromec','FG-AN-90-LC-2TNA','Em linha – disponível','Código mudou'),
 ('FM-EL-90-LC-2LCA','Forno Elétrico Luce 125L','Elettromec','FM-EL-90-LC-2TNA','Em linha – disponível','Código mudou'),
 ('GA-6S-14-VT-2TNA','Gaveta Aquecida Vetro 14cm','Elettromec','GA-6S-14-XX-2TNA (inox) ou GA-6S-14-BM-2TNA (black matte)','Em linha – indisponível','Vetro saiu; outros acabamentos'),
 ('BC-BI-86-XV-2ATB','Cervejeira 86L Connect 220V','Elettromec','BC-BI-88-XV-2VPA (88L)','Em linha – disponível','Modelo novo de 88L'),
 ('FB-BI-86-XV-2ATB','Frigobar 86L Connect 220V','Elettromec','FB-BI-88-XV-2VPA (88L)','Em linha – disponível','Modelo novo de 88L'),
 ('BC-BI-135-XV-2ATF','Cervejeira 135L Connect 220V','Elettromec','BC-BI-145-XV-2VPA (dir.) / 2VPB (esq.)','Em linha – disponível','Modelo novo de 145L'),
 ('BC-BI-135-XV-1ATF','Cervejeira 135L Connect 127V','Elettromec','BC-BI-145-XV-2VPA / 2VPB (só 220V)','Em linha – disponível','Novo só existe em 220V'),
 ('FB-BI-135-XV-2ATC','Frigobar 135L Connect 220V','Elettromec','FB-BI-145-XV-2VPA (145L)','Em linha – disponível','Modelo novo de 145L'),
 ('FB-BI-135-XV-1ATC','Frigobar 135L Connect 127V','Elettromec','FB-BI-145-XV-2VPA (só 220V)','Em linha – disponível','Novo só existe em 220V'),
]
deletar = [
 ('BBQ-3Q-65-XX-NHUA','Churrasqueira a Gás 65cm 3 Queimadores','Elettromec'),
 ('BBQ-4Q-30-XP-NHUA','Churrasqueira Professional Gás 30" 76cm','Elettromec'),
 ('DBQ-2Q-40-XX-NHUA','Dominó Churrasqueira Gás 2 Queimadores 40cm','Elettromec'),
 ('CFP-MLN-120-XX-2ATA','Coifa de Parede Milano 120cm','Elettromec'),
 ('CFP-RVN-90-XV-1ATB','Coifa de Parede Ravenna 90cm 127V','Elettromec'),
 ('CFP-RVN-90-XV-2ATB','Coifa de Parede Ravenna 90cm 220V','Elettromec'),
 ('CKG-6Q-90-XQ-3ZEB','Cooktop a Gás Quadratto 6 Bocas 90cm','Elettromec'),
 ('DV-2Q-30-XQ-2ZEA','Dominó Elétrico Vitrocerâmico 2 Bocas','Elettromec'),
 ('CKI-5Q-90-CI-2XBA','Cooktop Indução 5 Bocas 90cm','Elettromec'),
 ('CV-2BI-26-XV-2ATA','Adega 26 Garrafas Dual Zone Connect','Elettromec'),
 ('CV-2BI-160-XV-2ATB','Adega 160 Garrafas Dual Zone Connect','Elettromec'),
 ('GR-BI-105-XX-2VIA','Gaveta Refrigerada 105L','Elettromec'),
 ('RF-FD-653-XX-2VSA','Refrigerador French Door 653L','Elettromec'),
 ('RF-MD-630-XX-2VSA','Refrigerador Multidoor 630L','Elettromec'),
 ('ECTGD1QM3A','Cooktop Dominó Evol 1 Queimador Tripla Chama','Evol'),
 ('ERDEF4362A','Refrigerador Side by Side Evol 436L','Evol'),
 ('S12PZ','Forno de Pizza Portátil Evol Parma 12"','Evol'),
 ('S16PZ','Forno de Pizza Portátil Evol Parma 16"','Evol'),
]
wb = load_workbook('Pente-fino-SKUs-Unikitchen.xlsx')
if 'Plano de ação' in wb.sheetnames: del wb['Plano de ação']
ws = wb.create_sheet('Plano de ação', 1)
cab = (Font(name=F, size=10, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1A1A1A'))
fn = Font(name=F, size=10); b = Border(bottom=Side(style='thin', color='D9D9D9'))
def bloco(lin, titulo, sub, cols, dados):
    ws.cell(row=lin, column=1, value=titulo).font = Font(name=F, size=12, bold=True)
    ws.cell(row=lin+1, column=1, value=sub).font = Font(name=F, size=9, italic=True, color='666666')
    for j, t in enumerate(cols, 1):
        c = ws.cell(row=lin+2, column=j, value=t); c.font, c.fill = cab
    for i, row in enumerate(dados, lin+3):
        for j, v in enumerate(list(row) + [''], 1):
            c = ws.cell(row=i, column=j, value=v); c.font = fn; c.border = b
            c.alignment = Alignment(wrap_text=True, vertical='top')
    return lin + 3 + len(dados) + 2
n = bloco(1, f'1. ATUALIZAR ({len(atual)}) – próxima etapa', 'Produto segue na marca com outro código/modelo: trocar código e ficha técnica no anúncio.',
          ['SKU no site', 'Produto', 'Marca', 'Código/modelo na lista da marca', 'Situação na lista', 'Motivo', 'Feito?'], atual)
bloco(n, f'2. DELETAR DE VEZ ({len(deletar)})', 'Não aparecem de forma alguma nas listas (nem código, nem modelo parecido). Já estão fora do site.',
      ['SKU no site', 'Produto', 'Marca', 'Feito?'], deletar)
for col, w in zip('ABCDEFG', [22, 42, 12, 46, 26, 34, 9]): ws.column_dimensions[col].width = w
wb.save('Pente-fino-SKUs-Unikitchen.xlsx')
print(len(atual), len(deletar))
