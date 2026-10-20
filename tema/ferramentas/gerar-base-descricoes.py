# Gera base-descricoes-produtos.docx com todas as descrições entregues (lidas do controle.md).
# Uso: cd tema/descricoes && python3 ../ferramentas/gerar-base-descricoes.py
import re, os
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, Cm
L=open('controle.md').read().split('\n')
rows=[]
for l in L:
    c=[x.strip() for x in l.split('|')[1:-1]]
    if len(c)==8 and c[0] not in('','x','SKU (= código do fabricante)') and not c[0].startswith('---') and '.html' in c[3]:
        rows.append(c)
def seo_fields(path):
    d={}
    if not os.path.exists(path): return d
    t=open(path).read().split('\n')
    for i,l in enumerate(t):
        for k in ('NOME DO PRODUTO','TAG TITLE','META DESCRIPTION'):
            if l.strip()==k and i+1<len(t): d[k]=t[i+1].strip()
    return d
class P(HTMLParser):
    def __init__(s,doc): super().__init__(); s.doc=doc; s.buf=[]; s.table=None; s.row=None; s.bold=False; s.ci=0
    def flush(s,style=None):
        if ''.join(x for x,_ in s.buf).strip():
            p=s.doc.add_paragraph(style=style)
            for t,b in s.buf: r=p.add_run(t); r.bold=b
        s.buf=[]
    def handle_starttag(s,tag,a):
        if tag in('strong','b'): s.bold=True
        elif tag=='table': s.table=s.doc.add_table(rows=0,cols=2); s.table.style='Light Grid Accent 1'
        elif tag=='tr' and s.table is not None: s.row=s.table.add_row().cells; s.ci=0
        elif tag in('th','td') and s.row is not None: s.buf=[]
    def handle_endtag(s,tag):
        if tag in('strong','b'): s.bold=False
        elif tag=='h2': s.flush('Heading 3')
        elif tag=='h3': s.flush('Heading 4')
        elif tag=='p': s.flush()
        elif tag in('th','td') and s.row is not None:
            txt=re.sub(r'\s+',' ',''.join(x for x,_ in s.buf)).strip()
            if s.ci<2:
                s.row[s.ci].text=txt
                if tag=='th':
                    for r in s.row[s.ci].paragraphs[0].runs: r.bold=True
            s.ci+=1; s.buf=[]
        elif tag=='table': s.table=None; s.row=None
    def handle_data(s,d):
        d=re.sub(r'\s+',' ',d)
        if d.strip() or s.buf: s.buf.append((d,s.bold))
doc=Document()
for m in ('left_margin','right_margin'): setattr(doc.sections[0],m,Cm(2))
doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10)
doc.add_heading('Unikitchen – Base de descrições de produtos',0)
doc.add_paragraph(f'{len(rows)} produtos, na ordem em que foram feitos. Para cada um: SKU, marca, nome, título SEO, meta description e a descrição completa. O HTML de cada descrição está no repositório (tema/descricoes).')
for i,r in enumerate(rows):
    f=r[3].strip('`'); s=r[4].strip('`').split(' ')[0].strip('`')
    if i: doc.add_page_break()
    doc.add_heading(re.sub(r'[*`]','',r[0])+' – '+r[2],1)
    d=seo_fields(s)
    for k,lab in (('NOME DO PRODUTO','Nome'),('TAG TITLE','Título SEO'),('META DESCRIPTION','Meta description')):
        p=doc.add_paragraph(); a=p.add_run(lab+': '); a.bold=True; p.add_run(d.get(k,'-'))
    p=doc.add_paragraph(); a=p.add_run('Situação: '); a.bold=True; p.add_run(re.sub(r'[*`]','',r[5])+' ('+r[6]+')')
    P(doc).feed(open(f).read())
doc.save('base-descricoes-produtos.docx'); print(len(rows))
